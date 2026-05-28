from __future__ import annotations

import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SRC_ROOT = PROJECT_ROOT / "src"
sys.path.insert(0, str(SRC_ROOT))

from ispa_daily_workflow.domain.pear.intake import (
    REPORT_SUSPENSION,
    REPORT_SUSPENSION_VS_OVERDUE,
    extract_canonical_landing_report,
    extract_landing_report,
)
from ispa_daily_workflow.schema import ValidationError


def _write_canonical_suspension_vs_overdue(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "School Name",
            "Last Name",
            "First Name",
            "Date of Birth",
            "ClientID",
            "Suspension Effective From",
            "Suspension Delete Reason",
            "Suspension Rescind Date",
            "Overdue Student",
        ]
    )
    sheet.append(
        [
            "ALPHA SCHOOL - 1001",
            "DOE",
            "JANE",
            "2012-01-01",
            "0000000001",
            "2026-02-24",
            "",
            "",
            "Y",
        ]
    )
    workbook.save(path)


def _write_no_data_suspension_vs_overdue(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["ReminderRecallSuspension_NeedSuspensionDeleted_Detail"])
    sheet.append(["Suspension Delete Reason = Contains Value = Suspension Deleted"])
    sheet.append(["Overdue Student = Blank = No Longer Overdue for ISPA Diseases"])
    sheet.append([])
    sheet.append(["No Data Available"])
    sheet.append([datetime(2026, 2, 27), 1, "5:06:13 AM"])
    workbook.save(path)


def _write_no_data_suspension(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["On Suspensions List", "Secondary School"])
    sheet.append([])
    sheet.append(["No Data Available"])
    sheet.append([])
    sheet.append([datetime(2026, 3, 11), "1 of 1", "5:22:01 AM"])
    workbook.save(path)


class TestPearCanonicalLanding(unittest.TestCase):
    def test_extract_canonical_landing_report(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "20260224_suspension_vs_overdue.xlsx"
            _write_canonical_suspension_vs_overdue(path)

            landing = extract_canonical_landing_report(path)
            self.assertEqual(landing.report_type, REPORT_SUSPENSION_VS_OVERDUE)
            self.assertEqual(landing.footer.report_date.isoformat(), "2026-02-24")
            self.assertEqual(landing.canonical_filename, path.name)
            self.assertEqual(landing.landing_frame.height, 1)
            self.assertEqual(landing.warnings, ())

    def test_extract_canonical_landing_requires_matching_filename(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "20260224_overdue_list_pear.xlsx"
            _write_canonical_suspension_vs_overdue(path)

            with self.assertRaisesRegex(ValidationError, "expected canonical filename"):
                extract_canonical_landing_report(path)

    def test_extract_landing_report_handles_suspension_vs_overdue_no_data(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "raw_no_data.xlsx"
            _write_no_data_suspension_vs_overdue(path)

            landing = extract_landing_report(path)
            self.assertEqual(landing.report_type, REPORT_SUSPENSION_VS_OVERDUE)
            self.assertEqual(landing.footer.report_date.isoformat(), "2026-02-27")
            self.assertEqual(
                landing.canonical_filename, "20260227_suspension_vs_overdue.xlsx"
            )
            self.assertEqual(landing.landing_frame.height, 0)
            self.assertEqual(len(landing.source_rows), 0)
            self.assertTrue(
                any(
                    "No Data Available marker detected" in warning
                    for warning in landing.warnings
                )
            )

    def test_extract_landing_report_handles_suspension_no_data(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "raw_secondary_suspension_no_data.xlsx"
            _write_no_data_suspension(path)

            landing = extract_landing_report(path)
            self.assertEqual(landing.report_type, REPORT_SUSPENSION)
            self.assertEqual(landing.footer.report_date.isoformat(), "2026-03-11")
            self.assertEqual(
                landing.canonical_filename, "20260311_suspension_list_secondary.xlsx"
            )
            self.assertEqual(landing.landing_frame.height, 0)
            self.assertEqual(len(landing.source_rows), 0)
            self.assertTrue(
                any(
                    "No Data Available marker detected; interpreted as an empty suspension_list report"
                    in warning
                    for warning in landing.warnings
                )
            )


if __name__ == "__main__":
    unittest.main()
