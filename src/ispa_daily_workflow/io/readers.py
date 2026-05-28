from __future__ import annotations

import csv
import json
from pathlib import Path

import polars as pl

from ispa_daily_workflow.schema import LoadedSchema


def read_header(path: Path, file_format: str | None = None) -> list[str]:
    suffix = (file_format or path.suffix.lstrip(".")).lower()
    if suffix in {"csv", "tsv"}:
        delimiter = "," if suffix == "csv" else "\t"
        with path.open(newline="", encoding="utf-8", errors="ignore") as handle:
            reader = csv.reader(handle, delimiter=delimiter)
            header = next(reader, [])
        return [str(value).strip() for value in header]

    if suffix == "json":
        with path.open("r", encoding="utf-8") as handle:
            payload = json.load(handle)
        if isinstance(payload, list) and payload and isinstance(payload[0], dict):
            return [str(k).strip() for k in payload[0]]
        if isinstance(payload, dict):
            return [str(k).strip() for k in payload]
        return []

    if suffix == "xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        workbook.close()
        return [str(value).strip() if value is not None else "" for value in row]

    if suffix == "xls":
        import xlrd

        book = xlrd.open_workbook(path)
        sheet = book.sheet_by_index(0)
        row = sheet.row_values(0) if sheet.nrows else []
        return [str(value).strip() for value in row]

    raise ValueError(f"Unsupported format for header read: {path}")


def read_dataframe(
    path: Path,
    file_format: str | None = None,
    *,
    schema_overrides: dict[str, pl.DataType] | None = None,
) -> pl.DataFrame:
    suffix = (file_format or path.suffix.lstrip(".")).lower()

    if suffix == "csv":
        return pl.read_csv(
            path, schema_overrides=schema_overrides, infer_schema_length=1000
        )
    if suffix == "tsv":
        return pl.read_csv(
            path,
            separator="\t",
            schema_overrides=schema_overrides,
            infer_schema_length=1000,
        )
    if suffix == "json":
        return pl.read_json(path)
    if suffix in {"xlsx", "xls"}:
        # Calamine is required for .xls support.
        frame = pl.read_excel(
            path,
            engine="calamine",
            sheet_id=0,
            schema_overrides=schema_overrides,
        )
        if isinstance(frame, dict):
            if not frame:
                raise ValueError(f"Workbook has no sheets: {path}")
            return next(iter(frame.values()))
        return frame

    raise ValueError(f"Unsupported format for dataframe read: {path}")


def _schema_type_to_polars_dtype(
    field_type: str,
    *,
    as_string: bool,
) -> pl.DataType:
    if as_string:
        return pl.String()

    normalized = str(field_type).strip().lower()
    if normalized == "string":
        return pl.String()
    if normalized == "boolean":
        return pl.Boolean()
    if normalized == "integer":
        return pl.Int64()
    if normalized == "number":
        return pl.Float64()
    if normalized == "date":
        return pl.Date()
    return pl.String()


def schema_overrides_from_schema(
    schema: LoadedSchema,
    *,
    as_string: bool = False,
) -> dict[str, pl.DataType]:
    return {
        name: _schema_type_to_polars_dtype(
            schema.type_map.get(name, "string"),
            as_string=as_string,
        )
        for name in schema.field_names
    }


def read_dataframe_with_schema(
    path: Path,
    schema: LoadedSchema,
    file_format: str | None = None,
    *,
    as_string: bool = False,
) -> pl.DataFrame:
    return read_dataframe(
        path,
        file_format=file_format,
        schema_overrides=schema_overrides_from_schema(schema, as_string=as_string),
    )
