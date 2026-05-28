from __future__ import annotations

from pathlib import Path
from typing import Any

import polars as pl
from openpyxl import load_workbook

from panorama_compliance.schema import ValidationError

from .constants import (
    EXPECTED_HEADERS,
    HEADER_ALIASES,
    PAGE_TOKEN_PATTERN,
    REPORT_ORDER,
)
from .constants import REPORT_SUSPENSION, REPORT_SUSPENSION_VS_OVERDUE
from .models import FooterMetadata, LandingReport
from .naming import canonical_filename, parse_canonical_report_date
from .primitives import (
    cell_text,
    extract_int,
    header_token,
    parse_date_required,
    parse_time_required,
    row_is_blank,
    summary_row,
    to_scalar,
)


def discover_input_files(input_dir: Path) -> list[Path]:
    return sorted(path for path in input_dir.glob("*.xlsx") if path.is_file())


def load_xlsx_rows(path: Path) -> list[list[Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def extract_landing_report(path: Path) -> LandingReport:
    rows = load_xlsx_rows(path)
    if not rows:
        raise ValidationError(f"{path.name}: workbook is empty")

    report_type = _detect_report_type(rows)
    expected_headers = EXPECTED_HEADERS[report_type]
    no_data_fallback = False
    try:
        header_index, header_map = _find_header_row(rows, HEADER_ALIASES[report_type])
    except ValidationError:
        if not _is_empty_no_data_report(rows, report_type=report_type):
            raise
        no_data_fallback = True
        header_index = -1
        header_map = {}

    footer_min_row = header_index + 1 if header_index >= 0 else 0
    footer_index, footer = _extract_footer(rows, footer_min_row)
    data_end = (
        _resolve_data_end(rows, header_index + 1, footer_index)
        if header_index >= 0
        else footer_index
    )

    source_rows: list[dict[str, Any]] = []
    landing_rows: list[list[Any]] = []
    if header_index >= 0:
        for row in rows[header_index + 1 : data_end]:
            if row_is_blank(row):
                continue
            source_record: dict[str, Any] = {}
            landing_record: list[Any] = []
            for column in expected_headers:
                index = header_map[column]
                value = row[index] if index < len(row) else None
                source_record[column] = value
                landing_record.append(to_scalar(value))
            source_rows.append(source_record)
            landing_rows.append(landing_record)

    landing_frame = pl.DataFrame(
        landing_rows,
        schema=expected_headers,
        orient="row",
    )
    title_stop_row = header_index if header_index >= 0 else footer_index
    title_text = collect_title_text(rows, title_stop_row)
    filename = canonical_filename(
        report_type=report_type,
        title_text=title_text,
        report_date=footer.report_date,
        source_name=path.name,
    )

    warnings: list[str] = []
    if no_data_fallback:
        warnings.append(
            "No Data Available marker detected; interpreted as an empty "
            f"{report_type} report"
        )
    if footer.page_token is None:
        warnings.append("Footer page token was missing")
    elif footer.page_token.strip().lower() != "1 of 1":
        warnings.append(
            f"Footer page token was {footer.page_token!r}, expected '1 of 1'"
        )
    if footer.reported_total_count is not None and footer.reported_total_count != len(
        source_rows
    ):
        warnings.append(
            "Footer/summary reported_total_count did not match extracted row count"
        )

    return LandingReport(
        path=path,
        report_type=report_type,
        title_text=title_text,
        footer=footer,
        data_start_row=header_index + 1 if header_index >= 0 else footer_index,
        data_end_row=data_end,
        landing_frame=landing_frame,
        source_rows=source_rows,
        canonical_filename=filename,
        warnings=tuple(warnings),
    )


def extract_canonical_landing_report(path: Path) -> LandingReport:
    rows = load_xlsx_rows(path)
    if not rows:
        raise ValidationError(f"{path.name}: workbook is empty")

    report_type = _detect_canonical_landing_report_type(rows)
    report_date = parse_canonical_report_date(path, report_type=report_type)
    expected_headers = EXPECTED_HEADERS[report_type]
    header_index, header_map = _find_header_row(rows, HEADER_ALIASES[report_type])

    source_rows: list[dict[str, Any]] = []
    landing_rows: list[list[Any]] = []
    for row in rows[header_index + 1 :]:
        if row_is_blank(row):
            continue
        source_record: dict[str, Any] = {}
        landing_record: list[Any] = []
        for column in expected_headers:
            index = header_map[column]
            value = row[index] if index < len(row) else None
            source_record[column] = value
            landing_record.append(to_scalar(value))
        source_rows.append(source_record)
        landing_rows.append(landing_record)

    landing_frame = pl.DataFrame(
        landing_rows,
        schema=expected_headers,
        orient="row",
    )
    return LandingReport(
        path=path,
        report_type=report_type,
        title_text="",
        footer=FooterMetadata(
            report_date=report_date,
            report_time="00:00:00",
            page_token=None,
            reported_total_count=len(source_rows),
        ),
        data_start_row=header_index + 1,
        data_end_row=len(rows),
        landing_frame=landing_frame,
        source_rows=source_rows,
        canonical_filename=path.name,
        warnings=(),
    )


def detect_canonical_landing_report_type(rows: list[list[Any]]) -> str:
    return _detect_canonical_landing_report_type(rows)


def _detect_canonical_landing_report_type(rows: list[list[Any]]) -> str:
    matches: list[str] = []
    for report_type in REPORT_ORDER:
        try:
            _find_header_row(rows, HEADER_ALIASES[report_type])
            matches.append(report_type)
        except ValidationError:
            continue

    if len(matches) == 1:
        return matches[0]
    if not matches:
        raise ValidationError(
            "Unable to classify canonical PEAR landing report type from workbook headers"
        )
    raise ValidationError(
        f"Ambiguous canonical PEAR landing report type detection: {matches}"
    )


def detect_report_type(rows: list[list[Any]]) -> str:
    return _detect_report_type(rows)


def _detect_report_type(rows: list[list[Any]]) -> str:
    title_blob = collect_title_text(rows, min(6, len(rows))).replace(" ", "").lower()
    if "needsuspensiondeleted" in title_blob or "suspensionvsoverdue" in title_blob:
        return REPORT_SUSPENSION_VS_OVERDUE
    if "onsuspensionslist" in title_blob:
        return "suspension_list"
    if "forecastqueryoverdue" in title_blob or "remainingoverdue" in title_blob:
        return "overdue_summary"

    matches: list[str] = []
    for report_type in REPORT_ORDER:
        try:
            _find_header_row(rows, HEADER_ALIASES[report_type])
            matches.append(report_type)
        except ValidationError:
            continue

    if len(matches) == 1:
        return matches[0]
    if not matches:
        raise ValidationError(
            "Unable to classify PEAR report type from workbook layout"
        )
    raise ValidationError(f"Ambiguous PEAR report type detection: {matches}")


def find_header_row(
    rows: list[list[Any]], aliases: dict[str, list[str]]
) -> tuple[int, dict[str, int]]:
    return _find_header_row(rows, aliases)


def _find_header_row(
    rows: list[list[Any]], aliases: dict[str, list[str]]
) -> tuple[int, dict[str, int]]:
    search_limit = min(len(rows), 40)
    for row_index in range(search_limit):
        row = rows[row_index]
        token_index: dict[str, int] = {}
        for cell_index, cell_value in enumerate(row):
            text = cell_text(cell_value)
            if not text:
                continue
            token = header_token(text)
            if token and token not in token_index:
                token_index[token] = cell_index

        header_map: dict[str, int] = {}
        missing = False
        for canonical, allowed in aliases.items():
            matched_index: int | None = None
            for alias in allowed:
                alias_index = token_index.get(header_token(alias))
                if alias_index is not None:
                    matched_index = alias_index
                    break
            if matched_index is None:
                missing = True
                break
            header_map[canonical] = matched_index
        if not missing:
            return row_index, header_map
    raise ValidationError("Unable to locate expected header row")


def _is_empty_no_data_report(
    rows: list[list[Any]],
    *,
    report_type: str,
) -> bool:
    if report_type not in {REPORT_SUSPENSION, REPORT_SUSPENSION_VS_OVERDUE}:
        return False
    search_limit = min(len(rows), 40)
    for row in rows[:search_limit]:
        for value in row:
            token = header_token(cell_text(value))
            if "nodataavailable" in token:
                return True
    return False


def _extract_footer(
    rows: list[list[Any]],
    min_row: int,
) -> tuple[int, FooterMetadata]:
    nonblank_index = len(rows) - 1
    while nonblank_index >= min_row and row_is_blank(rows[nonblank_index]):
        nonblank_index -= 1
    if nonblank_index < min_row:
        raise ValidationError("Unable to locate footer metadata row")

    footer_row = rows[nonblank_index]
    nonblank_values = [value for value in footer_row if cell_text(value)]
    if not nonblank_values:
        raise ValidationError("Footer metadata row was blank")

    report_date = parse_date_required(
        nonblank_values[0], Path("."), "footer report date"
    )
    report_time = parse_time_required(
        nonblank_values[-1], Path("."), "footer report time"
    )

    page_token: str | None = None
    for value in nonblank_values:
        text = cell_text(value)
        if PAGE_TOKEN_PATTERN.match(text):
            page_token = text
            break

    reported_total_count: int | None = None
    summary_index = nonblank_index - 1
    while summary_index >= min_row and row_is_blank(rows[summary_index]):
        summary_index -= 1
    if summary_index >= min_row and summary_row(rows[summary_index]):
        reported_total_count = extract_int(rows[summary_index])

    return nonblank_index, FooterMetadata(
        report_date=report_date,
        report_time=report_time,
        page_token=page_token,
        reported_total_count=reported_total_count,
    )


def _resolve_data_end(rows: list[list[Any]], start: int, footer_index: int) -> int:
    end = footer_index
    while end > start and row_is_blank(rows[end - 1]):
        end -= 1
    if end > start and summary_row(rows[end - 1]):
        end -= 1
    while end > start and row_is_blank(rows[end - 1]):
        end -= 1
    return end


def collect_title_text(rows: list[list[Any]], stop_row: int) -> str:
    limit = stop_row if isinstance(stop_row, int) else len(rows)
    title_chunks: list[str] = []
    for row in rows[:limit]:
        for value in row:
            text = cell_text(value)
            if text:
                title_chunks.append(text)
    return " ".join(title_chunks).strip()
