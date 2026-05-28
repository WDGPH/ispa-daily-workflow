#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import math
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable

import openpyxl
import polars as pl
import xlrd


SUPPORTED_SUFFIXES = {".csv", ".json", ".parquet", ".xls", ".xlsx"}
DEFAULT_PATTERNS = ("*.parquet", "*.csv", "*.xlsx", "*.xls")


@dataclass
class ColumnProfile:
    name: str
    dtype: str | None = None
    profiled_rows: int = 0
    null_count: int = 0
    non_null_count: int = 0
    observed_types: set[str] = field(default_factory=set)
    distinct_limit: int = 10000
    distinct_values: set[str] = field(default_factory=set)
    distinct_overflow: bool = False
    min_length: int | None = None
    max_length: int | None = None
    total_length: int = 0
    include_ranges: bool = False
    min_value: Any = None
    max_value: Any = None
    include_examples: bool = False
    example_limit: int = 3
    examples: list[Any] = field(default_factory=list)

    def observe(self, value: Any) -> None:
        self.profiled_rows += 1
        if _is_null(value):
            self.null_count += 1
            return

        self.non_null_count += 1
        self.observed_types.add(_type_label(value))
        if not self.distinct_overflow:
            self.distinct_values.add(_stable_scalar(value))
            if len(self.distinct_values) > self.distinct_limit:
                self.distinct_values.clear()
                self.distinct_overflow = True

        if isinstance(value, str):
            length = len(value)
            self.min_length = (
                length if self.min_length is None else min(self.min_length, length)
            )
            self.max_length = (
                length if self.max_length is None else max(self.max_length, length)
            )
            self.total_length += length

        if self.include_ranges and _is_range_scalar(value):
            try:
                if self.min_value is None or value < self.min_value:
                    self.min_value = value
                if self.max_value is None or value > self.max_value:
                    self.max_value = value
            except TypeError:
                pass

        if self.include_examples and len(self.examples) < self.example_limit:
            encoded = _json_scalar(value)
            if encoded not in self.examples:
                self.examples.append(encoded)

    def as_dict(self) -> dict[str, Any]:
        distinct_count: int | str
        if self.distinct_overflow:
            distinct_count = f">{self.distinct_limit}"
        else:
            distinct_count = len(self.distinct_values)
        dtype = self.dtype or "/".join(sorted(self.observed_types)) or "unknown"
        profile: dict[str, Any] = {
            "name": self.name,
            "dtype": dtype,
            "profiled_null_count": self.null_count,
            "profiled_null_pct": _pct(self.null_count, self.profiled_rows),
            "profiled_unique_count": distinct_count,
            "observed_types": sorted(self.observed_types),
        }
        if self.min_length is not None:
            profile["min_length"] = self.min_length
            profile["max_length"] = self.max_length
            profile["mean_length"] = round(
                self.total_length / max(self.non_null_count, 1),
                2,
            )
        if self.include_ranges and self.min_value is not None:
            profile["min_value"] = _json_scalar(self.min_value)
            profile["max_value"] = _json_scalar(self.max_value)
        if self.include_examples:
            profile["examples"] = self.examples
        return profile


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Profile local CSV, Parquet, Excel, and explicit JSON files for "
            "fixture design. By default, the profile emits shapes, column "
            "names, types, null rates, cardinality counts, and string lengths, "
            "but no cell values."
        )
    )
    parser.add_argument(
        "paths",
        nargs="+",
        type=Path,
        help="Files or directories to profile.",
    )
    parser.add_argument(
        "--pattern",
        action="append",
        default=None,
        help="Glob pattern for directory inputs. Repeat as needed.",
    )
    parser.add_argument(
        "--recursive",
        action="store_true",
        help="Search directory inputs recursively.",
    )
    parser.add_argument(
        "--row-limit",
        type=int,
        default=50000,
        help="Rows per table to profile for column stats. Use 0 for all rows.",
    )
    parser.add_argument(
        "--distinct-limit",
        type=int,
        default=10000,
        help="Maximum distinct values tracked per column before reporting >limit.",
    )
    parser.add_argument(
        "--include-ranges",
        action="store_true",
        help="Include numeric/date min and max values in the profile.",
    )
    parser.add_argument(
        "--include-examples",
        action="store_true",
        help="Include up to three example cell values per column.",
    )
    parser.add_argument(
        "--format",
        choices=("json", "markdown"),
        default="json",
        help="Output format.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=None,
        help="Optional output path. Defaults to stdout.",
    )
    return parser.parse_args(argv)


def collect_files(
    paths: Iterable[Path],
    *,
    patterns: Iterable[str] = DEFAULT_PATTERNS,
    recursive: bool = False,
) -> list[Path]:
    files: list[Path] = []
    seen: set[Path] = set()
    for path in paths:
        if path.is_file():
            candidates = [path]
        elif path.is_dir():
            candidates = []
            for pattern in patterns:
                matches = path.rglob(pattern) if recursive else path.glob(pattern)
                candidates.extend(matches)
        else:
            continue

        for candidate in candidates:
            if candidate.name.startswith("~$"):
                continue
            if candidate.suffix.lower() not in SUPPORTED_SUFFIXES:
                continue
            resolved = candidate.resolve()
            if resolved not in seen:
                seen.add(resolved)
                files.append(candidate)
    return sorted(files)


def build_profile(
    paths: Iterable[Path],
    *,
    patterns: Iterable[str] = DEFAULT_PATTERNS,
    recursive: bool = False,
    row_limit: int = 50000,
    distinct_limit: int = 10000,
    include_ranges: bool = False,
    include_examples: bool = False,
) -> dict[str, Any]:
    files = collect_files(paths, patterns=patterns, recursive=recursive)
    profile: dict[str, Any] = {
        "generated_at": datetime.now(timezone.utc)
        .replace(microsecond=0)
        .isoformat()
        .replace("+00:00", "Z"),
        "privacy": (
            "No cell examples are emitted unless include_examples is true. "
            "Numeric/date ranges are emitted only when include_ranges is true."
        ),
        "row_limit": row_limit,
        "files": [],
        "errors": [],
    }
    for path in files:
        try:
            profile["files"].append(
                _profile_file(
                    path,
                    row_limit=row_limit,
                    distinct_limit=distinct_limit,
                    include_ranges=include_ranges,
                    include_examples=include_examples,
                )
            )
        except Exception as exc:  # pragma: no cover - defensive operational path
            profile["errors"].append({"path": str(path), "error": str(exc)})
    return profile


def _profile_file(
    path: Path,
    *,
    row_limit: int,
    distinct_limit: int,
    include_ranges: bool,
    include_examples: bool,
) -> dict[str, Any]:
    suffix = path.suffix.lower()
    if suffix == ".parquet":
        tables = [
            _profile_polars_table(
                pl.read_parquet(path),
                table_name="data",
                row_limit=row_limit,
                distinct_limit=distinct_limit,
                include_ranges=include_ranges,
                include_examples=include_examples,
            )
        ]
    elif suffix == ".csv":
        tables = [
            _profile_polars_table(
                pl.read_csv(path, infer_schema_length=1000, ignore_errors=True),
                table_name="data",
                row_limit=row_limit,
                distinct_limit=distinct_limit,
                include_ranges=include_ranges,
                include_examples=include_examples,
            )
        ]
    elif suffix == ".json":
        tables = _profile_json(
            path,
            row_limit=row_limit,
            distinct_limit=distinct_limit,
            include_ranges=include_ranges,
            include_examples=include_examples,
        )
    elif suffix == ".xlsx":
        tables = _profile_xlsx(
            path,
            row_limit=row_limit,
            distinct_limit=distinct_limit,
            include_ranges=include_ranges,
            include_examples=include_examples,
        )
    elif suffix == ".xls":
        tables = _profile_xls(
            path,
            row_limit=row_limit,
            distinct_limit=distinct_limit,
            include_ranges=include_ranges,
            include_examples=include_examples,
        )
    else:  # pragma: no cover - collect_files filters this
        raise ValueError(f"unsupported file type: {path.suffix}")

    return {
        "path": str(path),
        "format": suffix.lstrip("."),
        "size_bytes": path.stat().st_size,
        "tables": tables,
    }


def _profile_polars_table(
    frame: pl.DataFrame,
    *,
    table_name: str,
    row_limit: int,
    distinct_limit: int,
    include_ranges: bool,
    include_examples: bool,
) -> dict[str, Any]:
    profiled = frame if row_limit == 0 else frame.head(row_limit)
    columns = {
        name: ColumnProfile(
            name=name,
            dtype=str(dtype),
            distinct_limit=distinct_limit,
            include_ranges=include_ranges,
            include_examples=include_examples,
        )
        for name, dtype in frame.schema.items()
    }
    for row in profiled.iter_rows(named=True):
        for name, value in row.items():
            columns[name].observe(value)
    return {
        "name": table_name,
        "rows": frame.height,
        "profiled_rows": profiled.height,
        "column_count": len(frame.columns),
        "columns": [columns[name].as_dict() for name in frame.columns],
    }


def _profile_json(
    path: Path,
    *,
    row_limit: int,
    distinct_limit: int,
    include_ranges: bool,
    include_examples: bool,
) -> list[dict[str, Any]]:
    data = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(data, list):
        return [
            _profile_mapping_rows(
                "data",
                data,
                row_limit=row_limit,
                distinct_limit=distinct_limit,
                include_ranges=include_ranges,
                include_examples=include_examples,
            )
        ]
    if isinstance(data, dict):
        list_tables = {
            key: value
            for key, value in data.items()
            if isinstance(value, list) and value
        }
        if list_tables:
            return [
                _profile_mapping_rows(
                    key,
                    value,
                    row_limit=row_limit,
                    distinct_limit=distinct_limit,
                    include_ranges=include_ranges,
                    include_examples=include_examples,
                )
                for key, value in list_tables.items()
            ]
        return [
            _profile_mapping_rows(
                "data",
                [data],
                row_limit=row_limit,
                distinct_limit=distinct_limit,
                include_ranges=include_ranges,
                include_examples=include_examples,
            )
        ]
    return [
        _profile_mapping_rows(
            "data",
            [{"value": data}],
            row_limit=row_limit,
            distinct_limit=distinct_limit,
            include_ranges=include_ranges,
            include_examples=include_examples,
        )
    ]


def _profile_mapping_rows(
    table_name: str,
    rows: list[Any],
    *,
    row_limit: int,
    distinct_limit: int,
    include_ranges: bool,
    include_examples: bool,
) -> dict[str, Any]:
    mapping_rows = [row if isinstance(row, dict) else {"value": row} for row in rows]
    headers = sorted({key for row in mapping_rows for key in row})
    columns = [
        ColumnProfile(
            name=name,
            distinct_limit=distinct_limit,
            include_ranges=include_ranges,
            include_examples=include_examples,
        )
        for name in headers
    ]
    columns_by_name = {column.name: column for column in columns}
    profiled = mapping_rows if row_limit == 0 else mapping_rows[:row_limit]
    for row in profiled:
        for name in headers:
            columns_by_name[name].observe(row.get(name))
    return {
        "name": table_name,
        "rows": len(mapping_rows),
        "profiled_rows": len(profiled),
        "column_count": len(headers),
        "columns": [columns_by_name[name].as_dict() for name in headers],
    }


def _profile_xlsx(
    path: Path,
    *,
    row_limit: int,
    distinct_limit: int,
    include_ranges: bool,
    include_examples: bool,
) -> list[dict[str, Any]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return [
            _profile_rows(
                worksheet.title,
                worksheet.iter_rows(values_only=True),
                worksheet.max_row,
                row_limit=row_limit,
                distinct_limit=distinct_limit,
                include_ranges=include_ranges,
                include_examples=include_examples,
            )
            for worksheet in workbook.worksheets
        ]
    finally:
        workbook.close()


def _profile_xls(
    path: Path,
    *,
    row_limit: int,
    distinct_limit: int,
    include_ranges: bool,
    include_examples: bool,
) -> list[dict[str, Any]]:
    workbook = xlrd.open_workbook(path, on_demand=True)
    try:
        tables = []
        for sheet in workbook.sheets():
            rows = (sheet.row_values(index) for index in range(sheet.nrows))
            tables.append(
                _profile_rows(
                    sheet.name,
                    rows,
                    sheet.nrows,
                    row_limit=row_limit,
                    distinct_limit=distinct_limit,
                    include_ranges=include_ranges,
                    include_examples=include_examples,
                )
            )
        return tables
    finally:
        workbook.release_resources()


def _profile_rows(
    table_name: str,
    rows: Iterable[tuple[Any, ...] | list[Any]],
    total_rows_hint: int | None,
    *,
    row_limit: int,
    distinct_limit: int,
    include_ranges: bool,
    include_examples: bool,
) -> dict[str, Any]:
    iterator = iter(rows)
    header_row = next(iterator, ())
    headers = _normalize_headers(header_row)
    columns = [
        ColumnProfile(
            name=name,
            distinct_limit=distinct_limit,
            include_ranges=include_ranges,
            include_examples=include_examples,
        )
        for name in headers
    ]
    profiled_rows = 0
    total_rows = 0
    for row in iterator:
        if not any(not _is_null(value) for value in row):
            continue
        total_rows += 1
        if row_limit and profiled_rows >= row_limit:
            continue
        profiled_rows += 1
        padded = list(row) + [None] * max(0, len(columns) - len(row))
        for column, value in zip(columns, padded):
            column.observe(value)

    row_count = total_rows
    if total_rows_hint is not None:
        row_count = max(0, total_rows_hint - 1)
    return {
        "name": table_name,
        "rows": row_count,
        "profiled_rows": profiled_rows,
        "column_count": len(columns),
        "columns": [column.as_dict() for column in columns],
    }


def _normalize_headers(row: tuple[Any, ...] | list[Any]) -> list[str]:
    headers: list[str] = []
    seen: dict[str, int] = {}
    for index, value in enumerate(row, start=1):
        raw = str(value).strip() if value is not None else ""
        name = raw or f"column_{index}"
        count = seen.get(name, 0)
        seen[name] = count + 1
        if count:
            name = f"{name}_{count + 1}"
        headers.append(name)
    return headers


def _is_null(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, float) and math.isnan(value):
        return True
    return False


def _type_label(value: Any) -> str:
    if isinstance(value, bool):
        return "bool"
    if isinstance(value, int):
        return "int"
    if isinstance(value, float):
        return "float"
    if isinstance(value, Decimal):
        return "decimal"
    if isinstance(value, datetime):
        return "datetime"
    if isinstance(value, date):
        return "date"
    if isinstance(value, str):
        return "str"
    return type(value).__name__


def _stable_scalar(value: Any) -> str:
    encoded = _json_scalar(value)
    return json.dumps(encoded, sort_keys=True)


def _json_scalar(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    return value


def _is_range_scalar(value: Any) -> bool:
    if isinstance(value, bool):
        return False
    return isinstance(value, (int, float, Decimal, datetime, date))


def _pct(numerator: int, denominator: int) -> float:
    if denominator == 0:
        return 0.0
    return round((numerator / denominator) * 100, 2)


def render_markdown(profile: dict[str, Any]) -> str:
    lines = [
        "# Datafile Profile",
        "",
        f"Generated: {profile['generated_at']}",
        "",
        profile["privacy"],
        "",
    ]
    for file_profile in profile["files"]:
        lines.extend(
            [
                f"## {file_profile['path']}",
                "",
                f"Format: `{file_profile['format']}`",
                "",
            ]
        )
        for table in file_profile["tables"]:
            lines.extend(
                [
                    f"### {table['name']}",
                    "",
                    (
                        f"Rows: {table['rows']} "
                        f"(profiled {table['profiled_rows']}); "
                        f"columns: {table['column_count']}"
                    ),
                    "",
                    "| Column | Type | Nulls | Unique | Length |",
                    "|---|---:|---:|---:|---:|",
                ]
            )
            for column in table["columns"]:
                length = ""
                if "min_length" in column:
                    length = (
                        f"{column['min_length']}-{column['max_length']} "
                        f"(mean {column['mean_length']})"
                    )
                lines.append(
                    "| {name} | {dtype} | {nulls} ({pct}%) | {unique} | {length} |".format(
                        name=column["name"],
                        dtype=column["dtype"],
                        nulls=column["profiled_null_count"],
                        pct=column["profiled_null_pct"],
                        unique=column["profiled_unique_count"],
                        length=length,
                    )
                )
            lines.append("")
    if profile["errors"]:
        lines.extend(["## Errors", ""])
        for error in profile["errors"]:
            lines.append(f"- {error['path']}: {error['error']}")
        lines.append("")
    return "\n".join(lines)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    profile = build_profile(
        args.paths,
        patterns=args.pattern or DEFAULT_PATTERNS,
        recursive=args.recursive,
        row_limit=args.row_limit,
        distinct_limit=args.distinct_limit,
        include_ranges=args.include_ranges,
        include_examples=args.include_examples,
    )
    if args.format == "markdown":
        rendered = render_markdown(profile)
    else:
        rendered = json.dumps(profile, indent=2, default=_json_scalar) + "\n"
    if args.output is None:
        print(rendered, end="" if rendered.endswith("\n") else "\n")
    else:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(rendered, encoding="utf-8")
    return 1 if profile["errors"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
